"""XLSX parser (openpyxl, ``read_only=True, data_only=True, keep_links=False``).

* Each worksheet is one "page": a level-1 heading with the sheet name, then table blocks
  (first non-empty row = header, repeated per block). Chart sheets carry no text.
* Formulas are never evaluated: ``data_only`` returns the value cached by the authoring
  application (or nothing). External links are not loaded.
* Rows, columns and total cells are capped; trailing empty cells and empty rows are dropped.
* Hidden and very-hidden sheets are parsed too, but their text is also reported in the
  sheet heading's ``hidden`` field so the injection scanner sees it as a hidden channel.
"""

from __future__ import annotations

import io
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

import openpyxl

from docassist.ingestion.model import MAX_SHEET_NAMES, DocumentMetadata, Limits, ParsedDocument
from docassist.ingestion.parsers.base import (
    DocumentBuilder,
    ParseError,
    clip,
    one_line,
    render_row,
)

MAX_SHEET_NAME_CHARS = 200


def format_cell(value: Any) -> str:
    """Render a cached cell value the way a reader would see it."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds").removesuffix(" 00:00:00")
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return f"{value:.10g}"
    if isinstance(value, int | Decimal):
        return str(value)
    return one_line(str(value))


def _rows(
    sheet: Any, limits: Limits, builder: DocumentBuilder, budget: list[int]
) -> list[list[str]]:
    rows: list[list[str]] = []
    for raw in sheet.iter_rows(values_only=True, max_col=limits.max_sheet_cols):
        cells = [format_cell(value) for value in raw]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            continue
        rows.append(cells)
        budget[0] -= len(cells)
        if len(rows) >= limits.max_sheet_rows or budget[0] <= 0:
            builder.warn("rows_truncated")
            break
    return rows


def parse_xlsx(data: bytes, limits: Limits) -> ParsedDocument:
    builder = DocumentBuilder("xlsx", limits)
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True, keep_links=False, keep_vba=False
        )
    except Exception as exc:
        raise ParseError("parse_error", "unreadable xlsx") from exc
    names: list[str] = []
    try:
        budget = [limits.max_cells]
        for index, sheet in enumerate(workbook.worksheets, start=1):
            name = one_line(str(sheet.title))[:MAX_SHEET_NAME_CHARS] or f"Sheet {index}"
            names.append(name)
            builder.new_page(index)
            if builder.exhausted:
                break
            rows = _rows(sheet, limits, builder, budget)
            hidden = ""
            if getattr(sheet, "sheet_state", "visible") != "visible":
                builder.warn("hidden_sheet")
                hidden = "\n".join(render_row(row) for row in rows)[: limits.max_hidden_chars]
            builder.add_heading(name, 1, hidden=hidden)
            builder.add_table(rows)
            if budget[0] <= 0:
                break
        properties = workbook.properties
        metadata = DocumentMetadata(
            title=clip(properties.title),
            author=clip(properties.creator),
            subject=clip(properties.subject),
            created=properties.created.isoformat() if properties.created else None,
            modified=properties.modified.isoformat() if properties.modified else None,
            page_count=len(names),
            sheet_names=tuple(names[:MAX_SHEET_NAMES]),
            page_basis="sheets",
        )
    finally:
        workbook.close()
    document = builder.build(metadata)
    if not any(block.kind == "table" for page in document.pages for block in page.blocks):
        raise ParseError("empty_document")
    return document
