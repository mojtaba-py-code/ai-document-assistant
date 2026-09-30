"""Test helpers for the ingestion area: document generators and DB seeding.

Documents are generated in code (no binary fixtures): PDFs by a tiny writer below, DOCX via
python-docx, XLSX via openpyxl.
"""

from __future__ import annotations

import hashlib
import io
import uuid
import zlib
from dataclasses import dataclass
from typing import Any

# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
FAKE_JPEG = b"\xff\xd8\xff\xe0" + b"\x00\x10JFIF\x00" + b"\x00" * 64 + b"\xff\xd9"


@dataclass(frozen=True)
class PdfImage:
    width: int
    height: int
    data: bytes
    filter: str = "/FlateDecode"  # or "/DCTDecode"
    colorspace: str = "/DeviceGray"
    bits: int = 8


def _pdf_string(text: str) -> str:
    return "(" + text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ")"


def gray_image(width: int = 32, height: int = 16) -> PdfImage:
    pixels = bytes((x * 7 + y * 3) % 256 for y in range(height) for x in range(width))
    return PdfImage(width, height, zlib.compress(pixels))


def make_pdf(
    pages: list[list[str]],
    *,
    images: dict[int, PdfImage] | None = None,
    title: str | None = "Test PDF",
    author: str | None = "QA",
) -> bytes:
    """A valid PDF: one list of text lines per page (Helvetica), optional image per page."""
    images = images or {}
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog = add(b"")  # placeholder, filled below
    pages_obj = add(b"")
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    page_ids: list[int] = []
    for index, lines in enumerate(pages):
        ops = ["BT /F1 11 Tf 14 TL 72 740 Td"]
        for line in lines:
            ops.append(f"{_pdf_string(line)} Tj T*")
        ops.append("ET")
        resources = f"/Font << /F1 {font} 0 R >>"
        image = images.get(index)
        if image is not None:
            dict_ = (
                f"<< /Type /XObject /Subtype /Image /Width {image.width} /Height {image.height} "
                f"/ColorSpace {image.colorspace} /BitsPerComponent {image.bits} "
                f"/Filter {image.filter} /Length {len(image.data)} >>"
            )
            img = add(dict_.encode() + b"\nstream\n" + image.data + b"\nendstream")
            resources += f" /XObject << /Im1 {img} 0 R >>"
            ops.append("q 200 0 0 100 72 300 cm /Im1 Do Q")
        content = "\n".join(ops).encode("latin-1")
        stream = add(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
        page_ids.append(
            add(
                (
                    f"<< /Type /Page /Parent {pages_obj} 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << {resources} >> /Contents {stream} 0 R >>"
                ).encode()
            )
        )
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_obj} 0 R >>".encode()
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects[pages_obj - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    info_parts = []
    if title:
        info_parts.append(f"/Title {_pdf_string(title)}")
    if author:
        info_parts.append(f"/Author {_pdf_string(author)}")
    info_parts.append("/CreationDate (D:20260114093000Z)")
    info = add(("<< " + " ".join(info_parts) + " >>").encode())

    out = io.BytesIO()
    out.write(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R /Info {info} 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def encrypted_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf([["Top secret text"]]))))
    writer.encrypt("user-password", algorithm="AES-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


# --------------------------------------------------------------------------- #
# DOCX / XLSX
# --------------------------------------------------------------------------- #
def make_docx(
    *,
    hidden_run: str | None = None,
    page_break: bool = True,
    title: str = "Master Services Agreement",
    extra_paragraphs: list[str] | None = None,
) -> bytes:
    import docx
    from docx.shared import Pt, RGBColor

    document = docx.Document()
    document.core_properties.title = title
    document.core_properties.author = "Legal Team"
    document.add_heading(title, level=1)
    document.add_paragraph(
        "This Master Services Agreement is entered into by and between Acme Corporation, a "
        "Delaware corporation, and Globex LLC, effective as of 14 October 2026."
    )
    document.add_heading("Payment Terms", level=2)
    paragraph = document.add_paragraph("Invoices are payable Net 30 from the invoice date.")
    if hidden_run:
        run = paragraph.add_run(" " + hidden_run)
        run.font.hidden = True
    white = document.add_paragraph()
    white_run = white.add_run("Visible text.")
    white_run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    white_run.font.size = Pt(11)
    document.add_paragraph("First deliverable", style="List Bullet")
    document.add_paragraph("Second deliverable", style="List Bullet")
    table = document.add_table(rows=3, cols=3)
    for r, row in enumerate(
        [["Item", "Qty", "Price"], ["Widget", "2", "USD 100.00"], ["Gadget", "1", "USD 50.00"]]
    ):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    if page_break:
        document.add_page_break()
    document.add_heading("Term", level=2)
    document.add_paragraph("This Agreement expires on December 31, 2027 unless renewed.")
    for extra in extra_paragraphs or []:
        document.add_paragraph(extra)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_xlsx(sheets: dict[str, list[list[Any]]], *, hidden: tuple[str, ...] = ()) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
        if name in hidden:
            sheet.sheet_state = "hidden"
    workbook.properties.title = "Budget"
    workbook.properties.creator = "Finance"
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- #
# Realistic texts
# --------------------------------------------------------------------------- #
CONTRACT_TEXT = """# Supplier Agreement

This Supplier Agreement (the "Agreement") is made and entered into by and between Acme Corporation, a Delaware corporation ("Acme"), and Beta Supplies Ltd., an English company ("Supplier").

## 1. Term

The Effective Date of this Agreement is 2026-01-15. This Agreement shall remain in full force and effect until 31 December 2027 and expires on December 31, 2027 unless renewed. The renewal date is 01/12/2027.

## 2. Payment Terms

Supplier shall invoice monthly. Payment terms: Net 45. All invoices are payable within 45 days of receipt. The total contract value is USD 250,000.00.

## 3. Governing Law

This Agreement is governed by the laws of the State of New York. In witness whereof, the parties have executed this Agreement.
"""

INVOICE_TEXT = """INVOICE

Invoice Number: INV-2026-0042
Invoice Date: March 3, 2026
Due Date: 02/04/2026

Bill To: Globex LLC, 1 Main Street

Description | Qty | Unit Price | Amount
Consulting services | 10 | $150.00 | $1,500.00
Travel | 1 | $320.50 | $320.50

Subtotal: $1,820.50
Tax (VAT 20%): $364.10
Total Due: $2,184.60

Payment is due within 30 days. Please remit to accounts@globex.example.
"""


# --------------------------------------------------------------------------- #
# Database seeding (Document + DocumentVersion + encrypted blob)
# --------------------------------------------------------------------------- #
MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "txt": "text/plain",
    "md": "text/markdown",
}


@dataclass
class Seeded:
    org_id: uuid.UUID
    owner_id: uuid.UUID
    document_id: uuid.UUID
    version_id: uuid.UUID
    storage_key: str


async def seed_version(
    container: Any,
    storage: Any,
    *,
    org_id: uuid.UUID,
    owner_id: uuid.UUID,
    data: bytes,
    extension: str,
    classification: str = "INTERNAL",
    document_id: uuid.UUID | None = None,
    version_number: int = 1,
    doc_type: str = "other",
    doc_type_source: str = "auto",
    status: str = "processing",
    version_status: str = "uploaded",
    title: str = "Seeded document",
) -> Seeded:
    """Insert a document (unless ``document_id`` is given) and a version with a stored blob."""
    from docassist.db.models import Document, DocumentVersion
    from docassist.db.session import DbContext
    from docassist.documents.storage import new_object_key, storage_context

    version_id = uuid.uuid4()
    key = new_object_key(org_id)
    await storage.put_bytes(key, data, storage_context(org_id, "version", version_id))
    async with container.db.transaction(DbContext(org_id=org_id, user_id=owner_id)) as session:
        if document_id is None:
            document = Document(
                organization_id=org_id,
                owner_id=owner_id,
                title=title,
                classification=classification,
                doc_type=doc_type,
                doc_type_source=doc_type_source,
                status=status,
                version_count=1,
            )
            session.add(document)
            await session.flush()
            document_id = document.id
        session.add(
            DocumentVersion(
                id=version_id,
                organization_id=org_id,
                document_id=document_id,
                version_number=version_number,
                storage_key=key,
                original_filename=f"file.{extension}",
                extension=extension,
                declared_mime=MIME[extension],
                detected_mime=MIME[extension],
                size_bytes=len(data),
                sha256=hashlib.sha256(data).hexdigest(),
                status=version_status,
                created_by=owner_id,
            )
        )
    assert document_id is not None
    return Seeded(org_id, owner_id, document_id, version_id, key)


# --------------------------------------------------------------------------- #
# Test doubles
# --------------------------------------------------------------------------- #
class InlineSandbox:
    """Same parsers, JSON round trip and strict validation as the real sandbox - in process.

    Spawning an interpreter per document is what the real sandbox is for; most pipeline
    tests exercise everything *around* it, so they use this double to stay fast. The real
    sandbox has its own tests (``tests/unit/test_ingestion_sandbox.py``) and one end-to-end
    pipeline test.
    """

    def __init__(self, config: Any | None = None) -> None:
        from docassist.ingestion.sandbox import SandboxConfig

        self.config = config or SandboxConfig()
        self.calls: list[str] = []

    async def parse(self, fmt: str, data: bytes, *, ocr_images: bool = False) -> Any:
        import json

        from docassist.ingestion.model import parsed_document_from_json
        from docassist.ingestion.parsers import ParseError, parse_document
        from docassist.ingestion.sandbox import SandboxError

        self.calls.append(fmt)
        try:
            document = parse_document(fmt, data, self.config.limits, ocr_images=ocr_images)
        except ParseError as exc:
            raise SandboxError(exc.code) from exc
        payload = json.loads(json.dumps(document.to_json()))
        return parsed_document_from_json(payload, expected_format=fmt, limits=self.config.limits)


class RecordingEmbedder:
    """Deterministic embedder that records every text it receives (to prove none leaked)."""

    name = "recording"

    def __init__(
        self, *, dimensions: int = 1024, is_external: bool = True, fail: bool = False
    ) -> None:
        from docassist.embeddings.hashing import HashingEmbedder

        self.dimensions = dimensions
        self.model = "recording-v1"
        self.is_external = is_external
        self.fail = fail
        self.received: list[str] = []
        self._inner = HashingEmbedder(dimensions=dimensions)

    async def embed(self, texts: Any, *, kind: str) -> list[list[float]]:
        from docassist.embeddings.base import EmbeddingError

        if self.fail:
            raise EmbeddingError("provider down")
        self.received.extend(texts)
        return await self._inner.embed(texts, kind="document")

    async def aclose(self) -> None:
        return None
