"""The synthetic demo corpus: valid files, realistic facts, consistent with the upload rules."""

from __future__ import annotations

import io
import zipfile
from datetime import UTC, date, datetime, timedelta

import pytest
from pypdf import PdfReader

from docassist.authz.permissions import DEFAULT_CLEARANCE
from docassist.core.enums import Classification, DocumentType, Role
from docassist.core.redaction import PiiKind, contains_secret, find_pii
from docassist.demo.content import (
    INJECTION_PAYLOAD,
    ORGS,
    USERS,
    DemoDocument,
    build_documents,
    long_date,
)
from docassist.demo.pdf import PAGE_BREAK, build_pdf
from docassist.identity.schemas import normalize_email

TODAY = date(2026, 9, 30)
MAGIC = {"pdf": b"%PDF-", "docx": b"PK\x03\x04", "xlsx": b"PK\x03\x04"}


def _text_of(doc: DemoDocument) -> str:
    if doc.extension == "pdf":
        return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(doc.data)).pages)
    if doc.extension == "docx":
        from docx import Document as WordDocument

        word = WordDocument(io.BytesIO(doc.data))
        parts = [p.text for p in word.paragraphs]
        parts += [cell.text for table in word.tables for row in table.rows for cell in row.cells]
        return "\n".join(parts)
    if doc.extension == "xlsx":
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(doc.data), read_only=True, data_only=True)
        return "\n".join(
            " ".join(str(v) for v in row if v is not None)
            for sheet in book.worksheets
            for row in sheet.iter_rows(values_only=True)
        )
    return doc.data.decode("utf-8")


@pytest.fixture(scope="module")
def corpus() -> list[DemoDocument]:
    return build_documents(TODAY)


def test_every_file_is_a_valid_document_of_its_type(corpus: list[DemoDocument]) -> None:
    assert len(corpus) == 15
    assert len({d.key for d in corpus}) == len({d.title for d in corpus}) == len(corpus)
    for doc in corpus:
        if doc.extension in MAGIC:
            assert doc.data.startswith(MAGIC[doc.extension]), doc.key
        else:
            text = doc.data.decode("utf-8")
            assert "\x00" not in text
        text = _text_of(doc)
        missing = [fact for fact in doc.facts if fact not in text]
        assert not missing, (doc.key, missing)
    formats = {d.extension for d in corpus}
    assert formats == {"pdf", "docx", "xlsx", "csv", "md", "txt"}


def test_files_carry_no_active_content(corpus: list[DemoDocument]) -> None:
    for doc in corpus:
        if doc.extension == "pdf":
            for marker in (
                b"/JavaScript",
                b"/JS",
                b"/OpenAction",
                b"/Launch",
                b"/AA",
                b"/URI",
                b"/EmbeddedFile",
                b"/Encrypt",
                b"/XFA",
                b"/SubmitForm",
            ):
                assert marker not in doc.data, (doc.key, marker)
        elif doc.extension in {"docx", "xlsx"}:
            with zipfile.ZipFile(io.BytesIO(doc.data)) as archive:
                names = archive.namelist()
                assert not [
                    n for n in names if "vbaProject" in n or "activeX" in n or n.endswith(".bin")
                ]
                for name in names:
                    assert b'TargetMode="External"' not in archive.read(name), (doc.key, name)
        assert not contains_secret(_text_of(doc)), doc.key


def test_contract_expiry_dates_fall_in_the_next_30_to_200_days(corpus: list[DemoDocument]) -> None:
    contracts = [d for d in corpus if d.doc_type is DocumentType.CONTRACT]
    assert len(contracts) >= 5
    expiries = [d.expires_on for d in contracts]
    assert all(e is not None for e in expiries)
    assert len(set(expiries)) == len(expiries)  # all different
    for doc in contracts:
        assert doc.expires_on is not None
        assert TODAY + timedelta(days=30) <= doc.expires_on <= TODAY + timedelta(days=200)
        text = _text_of(doc)
        assert long_date(doc.expires_on) in text
        assert "Net " in text  # payment terms
    assert {d.org for d in contracts} == {"acme", "globex"}


def test_required_scenarios_are_present(corpus: list[DemoDocument]) -> None:
    by_key = {d.key: d for d in corpus}
    salary = by_key["acme.hr.salary_sheet"]
    assert (salary.extension, salary.classification, salary.department) == (
        "xlsx",
        Classification.RESTRICTED,
        "hr",
    )
    leave = by_key["acme.hr.annual_leave"]
    assert (leave.classification, leave.department) == (Classification.CONFIDENTIAL, "hr")
    auth = by_key["acme.engineering.auth_architecture"]
    assert (auth.classification, auth.department) == (Classification.INTERNAL, "engineering")
    assert by_key["acme.legal.nda_tailspin"].doc_type is DocumentType.LEGAL
    assert sum(1 for d in corpus if d.doc_type is DocumentType.INVOICE) >= 3
    assert sum(1 for d in corpus if d.org == "globex") >= 3
    carriers = [d.key for d in corpus if INJECTION_PAYLOAD in _text_of(d)]
    assert carriers == ["acme.finance.vendor_onboarding"]
    # the invoice carries bank details, so ingestion can suggest a higher classification
    invoice_pii = {m.kind for m in find_pii(_text_of(by_key["acme.invoice.northwind"]))}
    assert PiiKind.IBAN in invoice_pii


def test_uploaders_satisfy_the_upload_rules(corpus: list[DemoDocument]) -> None:
    users = {u.key: u for u in USERS}
    departments = {org.slug: {slug for slug, _ in org.departments} for org in ORGS}
    for doc in corpus:
        uploader = users[doc.uploader]
        assert uploader.org == doc.org, doc.key
        assert doc.department in departments[doc.org], doc.key
        # department must be one of the uploader's departments (org admins: any)
        assert uploader.role is Role.ORGANIZATION_ADMIN or doc.department in uploader.departments
        # classification may not exceed the uploader's clearance
        assert doc.classification.rank <= DEFAULT_CLEARANCE[uploader.role].rank, doc.key
        assert uploader.role is not Role.AUDITOR


def test_accounts_cover_every_role_in_each_tenant() -> None:
    emails = [normalize_email(u.email) for u in USERS]
    assert len(set(emails)) == len(emails)
    assert len({u.key for u in USERS}) == len(USERS)
    tenant_roles = {Role.ORGANIZATION_ADMIN, Role.DEPARTMENT_MANAGER, Role.EMPLOYEE, Role.AUDITOR}
    for org in ORGS:
        members = [u for u in USERS if u.org == org.slug]
        assert {u.role for u in members} == tenant_roles
        slugs = {slug for slug, _ in org.departments}
        for user in members:
            assert set(user.departments) <= slugs
            assert set(user.managed) <= set(user.departments)
            assert not user.managed or user.role is Role.DEPARTMENT_MANAGER
    platform = [u for u in USERS if u.org is None]
    assert [u.role for u in platform] == [Role.PLATFORM_ADMIN]
    assert [o.slug for o in ORGS] == ["acme", "globex"]
    assert {s for s, _ in ORGS[0].departments} == {"finance", "hr", "legal", "engineering"}
    assert {s for s, _ in ORGS[1].departments} == {"finance", "hr"}


def test_generation_is_deterministic_for_text_and_pdf(corpus: list[DemoDocument]) -> None:
    again = {d.key: d for d in build_documents(TODAY)}
    for doc in corpus:
        if doc.extension in {"pdf", "md", "txt", "csv"}:
            assert again[doc.key].data == doc.data, doc.key
    later = {d.key: d for d in build_documents(TODAY + timedelta(days=10))}
    northwind = later["acme.contract.northwind"]
    assert northwind.expires_on == TODAY + timedelta(days=55)


# --------------------------------------------------------------------------- #
# PDF writer
# --------------------------------------------------------------------------- #
def test_pdf_writer_escapes_wraps_and_paginates() -> None:
    tricky = (
        "Parens (a) and \\ backslash ) unbalanced ( and caf" + chr(0xE9) + " and " + chr(0x4E2D)
    )
    long_line = "word " * 60
    data = build_pdf(
        [tricky, long_line, PAGE_BREAK, *[f"line {i}" for i in range(60)]],
        title="T (x)",
        created=datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC),
    )
    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) == 3  # explicit break + 60 lines overflow the second page
    first = reader.pages[0].extract_text()
    assert "Parens (a) and \\ backslash ) unbalanced ( and caf" + chr(0xE9) + " and ?" in first
    assert first.count("word") == 60  # wrapped, nothing lost
    assert "line 0" in reader.pages[1].extract_text()
    assert "line 59" in reader.pages[2].extract_text()
    assert reader.metadata is not None and reader.metadata.title == "T (x)"
    assert data.rstrip().endswith(b"%%EOF")


def test_issued_credentials_never_render_the_password() -> None:
    from docassist.demo.seed import Credential, SeedReport

    secret = "S3cret-demo-password-value"
    credential = Credential("a@acme.example", "employee", "acme", secret)
    report = SeedReport(credentials=[credential])
    assert secret not in repr(credential)
    assert secret not in repr(report) and secret not in str(report)
